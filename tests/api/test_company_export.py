"""``/api/company/{name}/export``: one report, four formats, three-state fields kept three-state."""

from __future__ import annotations

import csv
import io
import json
from contextlib import nullcontext
from datetime import datetime, timezone

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from services.company_pages import PreparedSection
from services.exports import build_report, export_filename, render_report
from services.exports.protocol_report import NOT_DETERMINED, report_sections

OVERVIEW = {
    "company": "Example",
    "protocol_id": 7,
    "contract_count": 2,
    "contracts": [
        {
            "address": "0x" + "a" * 40,
            "name": "Vault",
            "chain": "ethereum",
            "role": "value_handler",
            "is_proxy": True,
            "proxy_type": "eip1967",
            "implementation": "0x" + "b" * 40,
            "owner": "0x" + "c" * 40,
            "source_verified": True,
            "is_pausable": True,
            "has_timelock": False,
            "upgrade_count": 3,
            "capabilities": ["pause", "upgradeable"],
            "total_usd": 1234.5,
        },
        {
            "address": "0x" + "d" * 40,
            "name": "Helper",
            "chain": "ethereum",
            "is_proxy": False,
            "source_verified": None,
            "is_pausable": None,
            "has_timelock": None,
        },
    ],
    "principals": [
        {
            "address": "0x" + "c" * 40,
            "chain": "ethereum",
            "type": "safe",
            "label": "Ops Safe",
            "details": {"owners": ["0x1", "0x2", "0x3"], "threshold": 2},
            "primary_for": ["0x" + "a" * 40],
            "co_controls": [],
            "chains": {"ethereum"},
            "controls_detail": [{"capabilities": ["upgrade", "pause"]}],
        },
        {
            "address": "0x" + "e" * 40,
            "chain": "ethereum",
            "type": "timelock",
            "label": "Timelock",
            "details": {"min_delay": 172800},
            "primary_for": [],
            "co_controls": ["0x" + "a" * 40],
        },
    ],
    "ownership_hierarchy": [],
    "fund_flows": [],
    "tvl": {"total_usd": 1234.5, "defillama_tvl": None},
    "member_analysis": {"members": 2, "analyzed": 2},
}

SCORE = {
    "company": "Example",
    "grade_state": "computed",
    "grade_lambda": 71.25,
    "grade_exposure": 0.4,
    "confidence_pct": 88.0,
    "perimeter_state": "complete",
    "model_version": "1.4.1-provisional",
    "computed_at": "2026-10-01T00:00:00+00:00",
    "trigger": "job",
    "findings": [
        {
            "capability": "upgrade",
            "principal_unit": "ethereum::0x" + "c" * 40,
            "raw_points": 12.0,
            "net_points_lambda": 12.0,
            "value_at_stake_usd": 1000.0,
            "n_functions": 1,
            "n_entities": 1,
            "example_functions": ["upgradeTo(address)"],
        }
    ],
}

COVERAGE = {
    "audit_count": 1,
    "coverage": [
        {
            "address": "0x" + "a" * 40,
            "chain": "ethereum",
            "contract_name": "Vault",
            "audit_count": 1,
            "last_audit": {"auditor": "Acme", "title": "Vault review", "date": "2026-01-01"},
        }
    ],
}

GENERATED = datetime(2026, 10, 13, 12, 0, tzinfo=timezone.utc)


def _report(score=SCORE, score_status="available", coverage=COVERAGE, overview=None):
    return build_report(
        company="Example",
        overview=overview or OVERVIEW,
        score=score,
        score_status=score_status,
        audit_coverage=coverage,
        generated_at=GENERATED,
    )


def test_report_summary_keeps_unread_fields_undetermined():
    summary = _report()["summary"]
    assert summary["contract_count"] == 2
    assert summary["proxy_count"] == 1
    assert summary["pausable"] == {"yes": 1, "no": 0, "not_determined": 1}
    assert summary["timelocked"] == {"yes": 0, "no": 1, "not_determined": 1}
    assert summary["multisig_count"] == 1
    assert summary["timelock_principal_count"] == 1
    assert summary["contracts_with_audit_coverage"] == 1


def test_sections_render_multisig_threshold_and_timelock_delay():
    principals = next(s for s in report_sections(_report()) if s.title == "Privileged principals")
    safe, timelock = principals.rows
    assert safe[4] == "2 of 3"
    assert safe[8] == "pause, upgrade"
    assert timelock[5] == "2d"


def test_contract_rows_never_render_unread_as_no():
    contracts = next(s for s in report_sections(_report()) if s.title == "Contracts")
    helper = next(row for row in contracts.rows if row[0] == "Helper")
    assert helper[4] == helper[8] == helper[9] == "Not determined"


def test_missing_score_says_why():
    report = _report(score=None, score_status="not_scored", coverage=None)
    score = next(s for s in report_sections(report) if s.title == "Score")
    assert "No score has been computed" in (score.note or "")
    audits = next(s for s in report_sections(report) if s.title == "Audit coverage")
    assert "could not be loaded" in (audits.note or "")
    assert report["summary"]["contracts_with_audit_coverage"] is None


def test_json_is_lossless_and_serializable():
    data = json.loads(render_report(_report(), "json"))
    assert data["protocol"] == {"name": "Example", "protocol_id": 7}
    assert data["principals"][0]["chains"] == ["ethereum"]
    assert data["score"]["findings"][0]["capability"] == "upgrade"


def test_csv_has_one_block_per_section():
    text = render_report(_report(), "csv").decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    titles = [r[0] for r in rows if r and r[0].startswith("# ")]
    assert titles == [
        "# Summary",
        "# Score",
        "# Score deductions",
        "# Contracts",
        "# Privileged principals",
        "# Audit coverage",
    ]
    assert any(r[:2] == ["Vault", "0x" + "a" * 40] for r in rows)


def test_pdf_is_readable():
    from pypdf import PdfReader

    content = render_report(_report(), "pdf")
    assert content.startswith(b"%PDF")
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
    assert "Example" in text and "Ops Safe" in text and "Not determined" in text


def test_docx_is_readable():
    from docx import Document

    document = Document(io.BytesIO(render_report(_report(), "docx")))
    headings = [p.text for p in document.paragraphs if p.style and (p.style.name or "").startswith("Heading")]
    assert "Contracts" in headings and "Privileged principals" in headings
    cells = {cell.text for table in document.tables for row in table.rows for cell in row.cells}
    assert "Vault" in cells and "2 of 3" in cells


def test_filename_is_slugged_and_stamped():
    assert export_filename("Lido / stETH", "pdf", GENERATED) == "psat-lido-steth-20261013T120000Z.pdf"


def _section(report, title):
    return next(s for s in report_sections(report) if s.title == title)


@pytest.fixture
def client(monkeypatch):
    from routers import company, deps

    monkeypatch.delenv("PSAT_PREPARED_COMPANY_PAGES", raising=False)
    monkeypatch.setattr(deps, "SessionLocal", lambda: nullcontext(None))
    monkeypatch.setattr(company, "eligible_company_names", lambda _session: {"Example": 7})
    monkeypatch.setattr(company, "build_company_overview", lambda _session, name: dict(OVERVIEW, company=name))
    monkeypatch.setattr(company, "company_score", lambda _name: SCORE)
    monkeypatch.setattr(company, "company_audit_coverage", lambda _name: COVERAGE)
    app = FastAPI()
    app.include_router(company.router)
    return TestClient(app)


@pytest.mark.parametrize(
    ("fmt", "media_type", "magic"),
    [
        ("json", "application/json", b"{"),
        ("csv", "text/csv", "﻿".encode()),
        ("pdf", "application/pdf", b"%PDF"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK"),
    ],
)
def test_route_serves_each_format_as_attachment(client, fmt, media_type, magic):
    response = client.get("/api/company/Example/export", params={"format": fmt})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(media_type)
    assert response.headers["content-disposition"].startswith('attachment; filename="psat-example-')
    assert response.headers["content-disposition"].endswith(f'.{fmt}"')
    assert response.content.startswith(magic)


def test_route_rejects_unknown_format(client):
    response = client.get("/api/company/Example/export", params={"format": "xlsx"})
    assert response.status_code == 400
    assert "json, csv, pdf, docx" in response.json()["detail"]


def test_route_404s_a_company_the_page_would_not_list(client, monkeypatch):
    """A protocol partway through its first analysis has a Protocol row but no page; it must not export."""
    from routers import company

    def build(_session, _name):
        raise AssertionError("an ineligible name must not reach the overview builder")

    monkeypatch.setattr(company, "build_company_overview", build)
    response = client.get("/api/company/Nope/export")
    assert response.status_code == 404
    assert response.json()["detail"] == "Company not found"


@pytest.mark.parametrize(("status", "expected"), [(404, "not_scored"), (503, "unavailable")])
def test_route_records_why_the_score_is_missing(client, monkeypatch, status, expected):
    from routers import company

    def no_score(_name):
        raise HTTPException(status_code=status, detail="x")

    monkeypatch.setattr(company, "company_score", no_score)
    body = client.get("/api/company/Example/export").json()
    assert body["score_status"] == expected
    assert body["score"] is None


@pytest.fixture
def prepared_client(client, monkeypatch):
    from routers import company

    monkeypatch.setenv("PSAT_PREPARED_COMPANY_PAGES", "1")

    def build(_session, _name):
        raise AssertionError("prepared mode must never build the overview live")

    monkeypatch.setattr(company, "build_company_overview", build)
    return client


def test_prepared_mode_exports_the_prepared_snapshot_without_building(prepared_client, monkeypatch):
    from routers import company

    prepared_at = datetime(2026, 10, 12, 9, 30, tzinfo=timezone.utc)
    sections = {
        "overview": PreparedSection(dict(OVERVIEW, tvl=None), prepared_at, False),
        "summary": PreparedSection({"tvl": {"total_usd": 99.0, "defillama_tvl": None}}, prepared_at, True),
    }
    monkeypatch.setattr(company, "prepared_section", lambda _s, _n, section="overview": sections[section])
    body = prepared_client.get("/api/company/Example/export").json()
    assert body["source"] == {"kind": "prepared-stale", "prepared_at": prepared_at.isoformat()}
    assert body["summary"]["tvl_usd"] == 99.0
    assert body["summary"]["contract_count"] == 2


@pytest.mark.parametrize(("eligible", "status"), [({"Example": 7}, 503), ({}, 404)])
def test_prepared_mode_answers_like_the_page_when_nothing_is_prepared(prepared_client, monkeypatch, eligible, status):
    from routers import company
    from services.aggregations.company_overview import jobs

    monkeypatch.setattr(company, "prepared_section", lambda *_a, **_k: None)
    monkeypatch.setattr(jobs, "eligible_company_names", lambda _session: eligible)
    response = prepared_client.get("/api/company/Example/export", params={"format": "pdf"})
    assert response.status_code == status
    if status == 503:
        assert response.json()["code"] == "company_preparing"
        assert response.headers["retry-after"] == "2"


UNREAD_OVERVIEW = {
    "protocol_id": 7,
    "contracts": [
        {"address": "0x" + "1" * 40, "name": "Unread", "chain": "ethereum"},
        {"address": "0x" + "2" * 40, "name": "Plain", "chain": "ethereum", "is_proxy": False, "capabilities": []},
    ],
    "principals": [
        {"address": "0x" + "3" * 40, "chain": "ethereum", "type": "safe", "label": None, "details": {}},
        {"address": "0x" + "4" * 40, "chain": "ethereum", "type": "safe", "details": {"threshold": 2}},
        {"address": "0x" + "5" * 40, "chain": "ethereum", "type": "timelock", "details": {}},
        {"address": "0x" + "6" * 40, "chain": "ethereum", "type": "eoa", "details": {}},
    ],
}


def test_unread_values_render_not_determined_and_inapplicable_ones_blank():
    report = _report(overview=UNREAD_OVERVIEW, coverage=None)
    unread, plain = _section(report, "Contracts").rows
    # Implementation, owner, upgrades, last upgrade, holdings, audits.
    for index in (6, 7, 10, 11, 13, 14):
        assert unread[index] == NOT_DETERMINED, index
    assert plain[6] == ""  # not a proxy: no implementation to report
    assert plain[12] == "None found"
    no_threshold, no_owners, timelock, eoa = _section(report, "Privileged principals").rows
    assert no_threshold[0] == NOT_DETERMINED
    assert no_threshold[4] == NOT_DETERMINED
    assert no_owners[4] == "2 of ? (signer count not determined)"
    assert timelock[5] == NOT_DETERMINED
    assert eoa[4] == eoa[5] == ""


def test_audit_counts_are_not_determined_when_coverage_could_not_be_loaded():
    contracts = _section(_report(coverage=None), "Contracts")
    assert {row[14] for row in contracts.rows} == {NOT_DETERMINED}
    listed = _section(_report(), "Contracts")
    assert {row[0]: row[14] for row in listed.rows} == {"Vault": "1", "Helper": NOT_DETERMINED}


def test_a_computed_score_with_a_null_figure_never_renders_blank():
    score = dict(SCORE, grade_lambda=None, grade_exposure=None, confidence_pct=None)
    pairs = dict(_section(_report(score=score), "Score").pairs)
    assert pairs["Grade (λ)"] == pairs["Exposure"] == pairs["Confidence"] == NOT_DETERMINED


def test_data_source_is_stated():
    live = build_report(
        company="Example",
        overview=OVERVIEW,
        score=SCORE,
        score_status="available",
        audit_coverage=COVERAGE,
        generated_at=GENERATED,
        source={"kind": "prepared", "prepared_at": "2026-10-12T09:30:00+00:00"},
    )
    assert dict(_section(live, "Summary").pairs)["Data source"] == "Prepared snapshot of 2026-10-12T09:30:00+00:00"
    assert dict(_section(_report(), "Summary").pairs)["Data source"] == NOT_DETERMINED


HOSTILE = [
    '=HYPERLINK("http://x.example","click")',
    "+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(1,1)",
    "\tTabbed",
    "\rReturn",
]


def _hostile_report():
    principals = [
        {"address": f"0x{i:040x}", "chain": "ethereum", "type": "eoa", "label": label}
        for i, label in enumerate(HOSTILE)
    ]
    coverage = {
        "audit_count": 1,
        "coverage": [
            {
                "address": "0x" + "a" * 40,
                "chain": "ethereum",
                "contract_name": "Vault",
                "audit_count": 1,
                "last_audit": {"auditor": "@evil", "title": "=1+1", "date": "2026-01-01"},
            }
        ],
    }
    return _report(overview=dict(OVERVIEW, principals=principals), coverage=coverage)


def test_csv_neutralises_formula_cells_from_untrusted_text():
    rows = list(csv.reader(io.StringIO(render_report(_hostile_report(), "csv").decode("utf-8-sig"))))
    cells = {cell for row in rows for cell in row}
    for value in HOSTILE + ["@evil", "=1+1"]:
        assert value not in cells, value
        assert "'" + value in cells, value
    # Plain negative numbers stay numeric.
    assert not any(cell.startswith("'-") and cell[2:].replace(".", "").isdigit() for cell in cells)


def test_csv_keeps_plain_negative_numbers():
    from services.exports.protocol_report import _csv_safe

    assert _csv_safe("-25.5") == "-25.5"
    assert _csv_safe("-1,000") == "-1,000"
    assert _csv_safe("-1+1") == "'-1+1"


@pytest.mark.parametrize("fmt", ["csv", "docx", "pdf"])
def test_control_characters_from_extracted_text_never_break_a_render(fmt):
    principals = [{"address": "0x" + "7" * 40, "chain": "ethereum", "type": "eoa", "label": "Ops\x00 Bell\x07\ufffe"}]
    content = render_report(_report(overview=dict(OVERVIEW, principals=principals, company="Ex\x01ample")), fmt)
    if fmt == "csv":
        assert "Ops Bell" in content.decode("utf-8-sig")
    if fmt == "docx":
        from docx import Document

        cells = {c.text for t in Document(io.BytesIO(content)).tables for r in t.rows for c in r.cells}
        assert "Ops Bell" in cells


def test_a_cell_taller_than_a_page_is_clipped_in_the_pdf():
    from pypdf import PdfReader

    title = "Very long audit title " * 4000
    coverage = {
        "audit_count": 1,
        "coverage": [
            {
                "address": "0x" + "a" * 40,
                "chain": "ethereum",
                "contract_name": "Vault",
                "audit_count": 1,
                "last_audit": {"auditor": "Acme", "title": title, "date": "2026-01-01"},
            }
        ],
    }
    report = _report(coverage=coverage)
    reader = PdfReader(io.BytesIO(render_report(report, "pdf")))
    assert "…" in "".join(page.extract_text() for page in reader.pages)
    # Only the PDF clips; the other formats keep the full text.
    assert title.strip() in render_report(report, "csv").decode("utf-8-sig")


def test_csv_omits_the_header_of_a_section_with_no_rows_to_show():
    rows = list(csv.reader(io.StringIO(render_report(_report(coverage=None), "csv").decode("utf-8-sig"))))
    at = rows.index(["# Audit coverage"])
    assert rows[at + 1] == ["Audit coverage could not be loaded when this export was generated."]
