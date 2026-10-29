"""Downloadable report of a protocol's latest results, in JSON, CSV, PDF or DOCX.

JSON carries the full overview, score document and audit coverage, losslessly. The other formats render a curated set
of sections from that same report, so CSV, PDF and DOCX cannot disagree with each other.

The tool publishes a value only when evidence proves it, so an unread value renders as "Not determined", never as a
default ("No", "0") or an empty cell, which a reader takes as "nothing here". A cell is left empty only where the column
doesn't apply to the row (a multisig threshold on a timelock).
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any

EXPORT_VERSION = 1

EXPORT_FORMATS: dict[str, str] = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

NOT_DETERMINED = "Not determined"
NOT_APPLICABLE = ""
# An empty list is the absence of a witness, not proof of absence.
NONE_FOUND = "None found"

# XML 1.0 forbids most C0 controls and U+FFFE/FFFF, so DOCX rejects them; csv rejects NUL. Audit text extracted from
# third-party PDFs carries them.
_UNWRITABLE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")
# Spreadsheets evaluate a cell starting with these as a formula (OWASP CSV injection).
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMBER = re.compile(r"-?\d[\d,]*(?:\.\d+)?%?")
# Rows of a PDF table can't split across pages; a cell taller than the frame raises LayoutError.
_PDF_CELL_MAX_LINES = 40
_PDF_FONT_SIZE = 6.5


class UnsupportedExportFormat(ValueError):
    pass


@dataclass
class Section:
    title: str
    # Exactly one of ``pairs`` (a key/value block) or ``columns`` + ``rows`` (a table).
    pairs: list[tuple[str, str]] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    note: str | None = None

    @property
    def is_table(self) -> bool:
        return bool(self.columns)


def _json_default(value: Any) -> Any:
    if isinstance(value, (set, frozenset, tuple)):
        return sorted(value, key=str)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def _plain(value: Any) -> Any:
    """Round-trip through JSON so the report holds only JSON types (sets, Decimals, datetimes from the DB)."""
    return json.loads(json.dumps(value, default=_json_default))


def _tri(value: Any) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return NOT_DETERMINED


def _clean(text: str) -> str:
    return _UNWRITABLE.sub("", text)


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _text(value: Any) -> str:
    if isinstance(value, str):
        value = _clean(value)
    if value is None or (isinstance(value, str) and not value.strip()):
        return NOT_DETERMINED
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    if isinstance(value, (list, tuple, set)):
        return _list(value)
    if isinstance(value, dict):
        return ", ".join(f"{k}: {_text(v)}" for k, v in value.items()) or NOT_DETERMINED
    return str(value)


def _list(values: Any) -> str:
    if values is None:
        return NOT_DETERMINED
    items = [_text(v) for v in values if v is not None and v != ""]
    return ", ".join(items) or NONE_FOUND


def _usd(value: Any) -> str:
    return f"${value:,.2f}" if _is_number(value) else NOT_DETERMINED


def _pct(value: Any) -> str:
    return f"{_text(value)}%" if _is_number(value) else NOT_DETERMINED


def _duration(seconds: Any) -> str:
    if isinstance(seconds, str) and seconds.isdigit():
        seconds = int(seconds)
    if not _is_number(seconds):
        return NOT_DETERMINED
    seconds = int(seconds)
    days, rem = divmod(seconds, 86400)
    hours, rem = divmod(rem, 3600)
    minutes = rem // 60
    parts = [f"{days}d" if days else "", f"{hours}h" if hours else "", f"{minutes}m" if minutes else ""]
    return " ".join(p for p in parts if p) or f"{seconds}s"


def _timelock_delay(details: dict[str, Any]) -> Any:
    for key in ("delay", "min_delay", "delay_seconds"):
        if details.get(key) is not None:
            return details[key]
    return None


def _multisig_threshold(details: dict[str, Any]) -> str:
    threshold = details.get("threshold")
    if isinstance(threshold, dict):
        threshold = threshold.get("m")
    if not _is_number(threshold):
        return NOT_DETERMINED
    owners = details.get("owners")
    if not isinstance(owners, list) or not owners:
        return f"{threshold} of ? (signer count not determined)"
    return f"{threshold} of {len(owners)}"


def _source_text(source: dict[str, Any] | None) -> str:
    kind = (source or {}).get("kind")
    prepared_at = (source or {}).get("prepared_at")
    if kind == "live":
        return "Built from the database when this file was generated"
    if kind in ("prepared", "prepared-stale") and prepared_at:
        stale = " (newer data was still being prepared)" if kind == "prepared-stale" else ""
        return f"Prepared snapshot of {prepared_at}{stale}"
    return NOT_DETERMINED


def _summary(
    overview: dict[str, Any], score: dict[str, Any] | None, audit_coverage: dict[str, Any] | None
) -> dict[str, Any]:
    contracts = overview.get("contracts") or []
    principals = overview.get("principals") or []
    current = {((c.get("chain") or "").lower(), (c.get("address") or "").lower()) for c in contracts}
    covered = [
        row
        for row in (audit_coverage or {}).get("coverage") or []
        if ((row.get("chain") or "").lower(), (row.get("address") or "").lower()) in current and row.get("audit_count")
    ]

    def tally(key: str) -> dict[str, int]:
        return {
            "yes": sum(1 for c in contracts if c.get(key) is True),
            "no": sum(1 for c in contracts if c.get(key) is False),
            "not_determined": sum(1 for c in contracts if c.get(key) is None),
        }

    tvl = overview.get("tvl") or {}
    return {
        "contract_count": len(contracts),
        "proxy_count": sum(1 for c in contracts if c.get("is_proxy")),
        "source_verified": tally("source_verified"),
        "pausable": tally("is_pausable"),
        "timelocked": tally("has_timelock"),
        "principal_count": len(principals),
        "multisig_count": sum(1 for p in principals if p.get("type") == "safe"),
        "timelock_principal_count": sum(1 for p in principals if p.get("type") == "timelock"),
        "eoa_principal_count": sum(1 for p in principals if p.get("type") == "eoa"),
        "audit_report_count": (audit_coverage or {}).get("audit_count"),
        "contracts_with_audit_coverage": len(covered) if audit_coverage is not None else None,
        "tvl_usd": tvl.get("total_usd"),
        "defillama_tvl_usd": tvl.get("defillama_tvl"),
        "grade_state": (score or {}).get("grade_state"),
        "score_finding_count": len((score or {}).get("findings") or []) if score else None,
    }


def build_report(
    *,
    company: str,
    overview: dict[str, Any],
    score: dict[str, Any] | None,
    score_status: str,
    audit_coverage: dict[str, Any] | None,
    generated_at: datetime,
    source: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """``score_status`` is ``available``, ``not_scored`` or ``unavailable`` (the document couldn't be read): a missing
    score must say which. ``source`` says where the overview came from: ``{"kind": "live" | "prepared" |
    "prepared-stale", "prepared_at": iso | None}``.
    """
    report = {
        "export_version": EXPORT_VERSION,
        "generated_at": generated_at.isoformat(),
        "source": source,
        "protocol": {"name": company, "protocol_id": overview.get("protocol_id")},
        "summary": _summary(overview, score, audit_coverage),
        "score_status": score_status,
        "score": score,
        "contracts": overview.get("contracts") or [],
        "principals": overview.get("principals") or [],
        "ownership_hierarchy": overview.get("ownership_hierarchy") or [],
        "fund_flows": overview.get("fund_flows") or [],
        "tvl": overview.get("tvl"),
        "member_analysis": overview.get("member_analysis"),
        "audit_coverage": audit_coverage,
    }
    return _plain(report)


def _score_section(report: dict[str, Any]) -> Section:
    score = report.get("score")
    status = report.get("score_status")
    if not score:
        note = (
            "No score has been computed for this protocol yet."
            if status == "not_scored"
            else "The score document could not be read when this export was generated."
        )
        return Section("Score", pairs=[("Status", status or NOT_DETERMINED)], note=note)
    determined = score.get("grade_state") == "computed"
    pairs = [
        ("Grade state", _text(score.get("grade_state"))),
        ("Grade (λ)", _text(score.get("grade_lambda")) if determined else NOT_DETERMINED),
        ("Exposure", _text(score.get("grade_exposure")) if determined else NOT_DETERMINED),
        ("Confidence", _pct(score.get("confidence_pct")) if determined else NOT_DETERMINED),
        ("Perimeter", _text(score.get("perimeter_state"))),
        ("Model version", _text(score.get("model_version"))),
        ("Computed at", _text(score.get("computed_at"))),
        ("Trigger", _text(score.get("trigger"))),
    ]
    return Section("Score", pairs=pairs)


def _findings_section(report: dict[str, Any]) -> Section:
    findings = (report.get("score") or {}).get("findings") or []
    rows = []
    for index, f in enumerate(findings, start=1):
        if not isinstance(f, dict):
            continue
        rows.append(
            [
                str(index),
                _text(f.get("capability")),
                _text(f.get("principal_unit")),
                _text(f.get("raw_points")),
                _text(f.get("net_points_lambda")),
                _usd(f.get("value_at_stake_usd")),
                _text(f.get("n_functions")),
                _text(f.get("n_entities")),
                _list((f.get("example_functions") or [])[:3] if "example_functions" in f else None),
            ]
        )
    return Section(
        "Score deductions",
        columns=[
            "#",
            "Capability",
            "Principal",
            "Raw points",
            "Net points (λ)",
            "Value at stake",
            "Functions",
            "Entities",
            "Example functions",
        ],
        rows=rows,
        note=None if rows else "No deductions.",
    )


def _contracts_section(report: dict[str, Any]) -> Section:
    loaded = report.get("audit_coverage")
    coverage = {
        ((row.get("chain") or "").lower(), (row.get("address") or "").lower()): row
        for row in (loaded or {}).get("coverage") or []
    }
    rows = []
    for c in report.get("contracts") or []:
        key = ((c.get("chain") or "").lower(), (c.get("address") or "").lower())
        # Unloaded coverage, or a contract it doesn't list, is unknown; "0 audits" would be a claim.
        audits = coverage.get(key, {}).get("audit_count") if loaded is not None else None
        proxy = _tri(c.get("is_proxy"))
        if c.get("is_proxy") and c.get("proxy_type"):
            proxy = f"Yes ({c['proxy_type']})"
        if c.get("is_proxy") is False:
            implementation = NOT_APPLICABLE
        elif c.get("is_proxy") is True:
            implementation = _text(c.get("implementation"))
        else:
            implementation = NOT_DETERMINED
        rows.append(
            [
                _text(c.get("name")),
                _text(c.get("address")),
                _text(c.get("chain")),
                _text(c.get("role")),
                _tri(c.get("source_verified")),
                proxy,
                implementation,
                _text(c.get("owner")),
                _tri(c.get("is_pausable")),
                _tri(c.get("has_timelock")),
                _text(c.get("upgrade_count")),
                _text(c.get("last_upgrade_timestamp")),
                _list(c.get("capabilities")),
                _usd(c.get("total_usd")),
                _text(audits),
            ]
        )
    return Section(
        "Contracts",
        columns=[
            "Name",
            "Address",
            "Chain",
            "Role",
            "Verified",
            "Proxy",
            "Implementation",
            "Owner",
            "Pausable",
            "Timelock",
            "Upgrades",
            "Last upgrade",
            "Capabilities",
            "Holdings (USD)",
            "Audits",
        ],
        rows=rows,
        note=None if rows else "No contracts mapped.",
    )


def _principals_section(report: dict[str, Any]) -> Section:
    rows = []
    for p in report.get("principals") or []:
        details = p.get("details") if isinstance(p.get("details"), dict) else {}
        controls = p.get("controls_detail") or []
        capabilities = sorted({cap for entry in controls for cap in entry.get("capabilities") or []})
        rows.append(
            [
                _text(p.get("label")),
                _text(p.get("address")),
                _text(p.get("chain")),
                _text(p.get("type")),
                _multisig_threshold(details) if p.get("type") == "safe" else NOT_APPLICABLE,
                _duration(_timelock_delay(details)) if p.get("type") == "timelock" else NOT_APPLICABLE,
                str(len(p.get("primary_for") or [])),
                str(len(p.get("co_controls") or [])),
                _list(capabilities),
            ]
        )
    return Section(
        "Privileged principals",
        columns=[
            "Label",
            "Address",
            "Chain",
            "Type",
            "Multisig threshold",
            "Timelock delay",
            "Controls",
            "Co-controls",
            "Capabilities",
        ],
        rows=rows,
        note=None if rows else "No privileged principals resolved.",
    )


def _audits_section(report: dict[str, Any]) -> Section:
    coverage = report.get("audit_coverage")
    if coverage is None:
        return Section("Audit coverage", note="Audit coverage could not be loaded when this export was generated.")
    rows = []
    for row in coverage.get("coverage") or []:
        last = row.get("last_audit")
        if isinstance(last, dict):
            latest = [_text(last.get("auditor")), _text(last.get("title")), _text(last.get("date"))]
        elif row.get("audit_count") == 0:
            latest = [NOT_APPLICABLE] * 3
        else:
            latest = [NOT_DETERMINED] * 3
        rows.append(
            [
                _text(row.get("contract_name")),
                _text(row.get("address")),
                _text(row.get("chain")),
                _text(row.get("audit_count")),
                *latest,
            ]
        )
    return Section(
        "Audit coverage",
        columns=["Contract", "Address", "Chain", "Audits", "Last auditor", "Last audit", "Last audit date"],
        rows=rows,
        note=None if rows else "No audit coverage on file.",
    )


def _summary_section(report: dict[str, Any]) -> Section:
    s = report.get("summary") or {}

    def tri_counts(key: str) -> str:
        t = s.get(key) or {}
        return f"{t.get('yes', 0)} yes · {t.get('no', 0)} no · {t.get('not_determined', 0)} not determined"

    pairs = [
        ("Protocol", _text((report.get("protocol") or {}).get("name"))),
        ("Generated at", _text(report.get("generated_at"))),
        ("Data source", _source_text(report.get("source"))),
        ("Contracts", _text(s.get("contract_count"))),
        ("Proxies", _text(s.get("proxy_count"))),
        ("Source verified", tri_counts("source_verified")),
        ("Pausable", tri_counts("pausable")),
        ("Timelocked", tri_counts("timelocked")),
        ("Privileged principals", _text(s.get("principal_count"))),
        ("Multisigs", _text(s.get("multisig_count"))),
        ("Timelock principals", _text(s.get("timelock_principal_count"))),
        ("EOA principals", _text(s.get("eoa_principal_count"))),
        ("Audit reports on file", _text(s.get("audit_report_count"))),
        ("Contracts with audit coverage", _text(s.get("contracts_with_audit_coverage"))),
        ("TVL (observed holdings)", _usd(s.get("tvl_usd"))),
        ("TVL (DefiLlama)", _usd(s.get("defillama_tvl_usd"))),
    ]
    return Section("Summary", pairs=pairs)


def _cleaned(section: Section) -> Section:
    return Section(
        _clean(section.title),
        pairs=[(_clean(k), _clean(v)) for k, v in section.pairs],
        columns=[_clean(c) for c in section.columns],
        rows=[[_clean(v) for v in row] for row in section.rows],
        note=_clean(section.note) if section.note else None,
    )


def report_sections(report: dict[str, Any]) -> list[Section]:
    """The readable sections every non-JSON format renders, with unwritable characters already removed."""
    sections = [
        _summary_section(report),
        _score_section(report),
        _findings_section(report),
        _contracts_section(report),
        _principals_section(report),
        _audits_section(report),
    ]
    return [_cleaned(section) for section in sections]


def _title(report: dict[str, Any]) -> str:
    return _clean(f"{(report.get('protocol') or {}).get('name') or 'Protocol'} — PSAT security report")


def render_json(report: dict[str, Any]) -> bytes:
    return json.dumps(report, indent=2, ensure_ascii=False).encode("utf-8")


def _csv_safe(value: str) -> str:
    """Prefix a would-be formula with ``'`` so a spreadsheet shows it as text; plain numbers stay numeric."""
    if value.startswith(_FORMULA_START) and not _PLAIN_NUMBER.fullmatch(value):
        return "'" + value
    return value


def render_csv(report: dict[str, Any]) -> bytes:
    """One file, one block per section: a ``# Title`` line, a header row, then rows, separated by a blank line."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")

    def write(row: list[str]) -> None:
        writer.writerow([_csv_safe(value) for value in row])

    for index, section in enumerate(report_sections(report)):
        if index:
            writer.writerow([])
        write([f"# {section.title}"])
        if section.is_table:
            write(section.columns)
            for row in section.rows:
                write(row)
        elif section.pairs:
            write(["Field", "Value"])
            for pair in section.pairs:
                write(list(pair))
        if section.note:
            write([section.note])
    # BOM so Excel reads UTF-8 (λ, —) instead of the system code page.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _column_widths(section: Section, total: float) -> list[float]:
    """Widths in proportion to each column's longest unbreakable word, so cells wrap at spaces rather than mid-word.

    Capped so a 42-character address can't starve the other columns; addresses wrap once.
    """
    weights = []
    for index, header in enumerate(section.columns):
        texts = [header, *(row[index] for row in section.rows)]
        longest = max((len(word) for text in texts for word in text.split()), default=0)
        weights.append(min(max(longest, 4), 22) + 2)
    scale = total / sum(weights)
    return [w * scale for w in weights]


def render_pdf(report: dict[str, Any]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=_PDF_FONT_SIZE, leading=8)
    head = cell.clone("head", fontName="Helvetica-Bold")
    # The built-in Helvetica is WinAnsi-encoded: no λ glyph.
    ascii_safe = lambda s: escape(s.replace("λ", "lambda"))  # noqa: E731

    def clipped(text: str, width: float) -> str:
        # Generous per-character estimate, so a clipped cell always fits; JSON, CSV and DOCX keep the full text.
        limit = max(1, int((width - 6) / (_PDF_FONT_SIZE * 0.45))) * _PDF_CELL_MAX_LINES
        return text if len(text) <= limit else text[: limit - 1] + "…"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=ascii_safe(_title(report)),
        author="PSAT",
    )
    story: list[Any] = [Paragraph(ascii_safe(_title(report)), styles["Title"])]
    grid = TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
    )
    for section in report_sections(report):
        story.append(Paragraph(ascii_safe(section.title), styles["Heading2"]))
        if section.is_table and section.rows:
            widths = _column_widths(section, doc.width)
            data = [[Paragraph(ascii_safe(c), head) for c in section.columns]]
            data += [
                [Paragraph(ascii_safe(clipped(v, w)), cell) for v, w in zip(row, widths, strict=True)]
                for row in section.rows
            ]
            table = Table(data, repeatRows=1, colWidths=widths)
            table.setStyle(grid)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0"))]))
            story.append(table)
        elif section.pairs:
            key_width, value_width = doc.width * 0.25, doc.width * 0.5
            data = [
                [Paragraph(ascii_safe(k), head), Paragraph(ascii_safe(clipped(v, value_width)), cell)]
                for k, v in section.pairs
            ]
            table = Table(data, colWidths=[key_width, value_width], hAlign="LEFT")
            table.setStyle(grid)
            story.append(table)
        if section.note:
            story.append(Paragraph(ascii_safe(section.note), styles["Italic"]))
        story.append(Spacer(1, 6 * mm))
    doc.build(story)
    return buffer.getvalue()


def render_docx(report: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt

    document = Document()
    page = document.sections[0]
    page.orientation = WD_ORIENT.LANDSCAPE
    page.page_width, page.page_height = Mm(297), Mm(210)
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(page, margin, Mm(12))
    document.core_properties.title = _title(report)
    document.core_properties.author = "PSAT"
    document.add_heading(_title(report), level=0)

    def add_table(rows: list[list[str]], *, header_row: bool, header_col: bool, size: float) -> None:
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for r, values in enumerate(rows):
            cells = table.add_row().cells
            for c, value in enumerate(values):
                run = cells[c].paragraphs[0].add_run(value)
                run.font.size = Pt(size)
                run.bold = (header_row and r == 0) or (header_col and c == 0)

    for section in report_sections(report):
        document.add_heading(section.title, level=1)
        if section.is_table and section.rows:
            add_table([section.columns, *section.rows], header_row=True, header_col=False, size=7)
        elif section.pairs:
            add_table([list(pair) for pair in section.pairs], header_row=False, header_col=True, size=9)
        if section.note:
            document.add_paragraph().add_run(section.note).italic = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_RENDERERS = {"json": render_json, "csv": render_csv, "pdf": render_pdf, "docx": render_docx}


def render_report(report: dict[str, Any], fmt: str) -> bytes:
    renderer = _RENDERERS.get(fmt)
    if renderer is None:
        raise UnsupportedExportFormat(fmt)
    return renderer(report)


def export_filename(company: str, fmt: str, generated_at: datetime) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", company).strip("-").lower() or "protocol"
    return f"psat-{slug}-{generated_at.strftime('%Y%m%dT%H%M%SZ')}.{fmt}"
